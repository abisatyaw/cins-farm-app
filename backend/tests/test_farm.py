from datetime import timedelta
from io import BytesIO

import pytest
from alembic.autogenerate import compare_metadata
from alembic.migration import MigrationContext
from app.config import now
from app.db import migrate
from conftest import NEW_PASSWORD, TEST_PASSWORD
from openpyxl import load_workbook
from sqlmodel import SQLModel


def create(client, kind, data):
    response = client.post(f"/api/records/{kind}", json=data)
    assert response.status_code == 201, response.text
    return response.json()


def sale(opt, **overrides):
    return {
        "date": str(now().date()),
        "product_id": opt("egg"),
        "unit_id": opt("kg"),
        "quantity": "10",
        "unit_price": 25000,
        "buyer": "Toko Sari",
        **overrides,
    }


def test_health_and_migration_drift(fresh):
    assert fresh.get("/health").json()["timezone"] == "Asia/Jakarta"
    engine = fresh.app.state.engine
    with engine.connect() as connection:
        assert compare_metadata(MigrationContext.configure(connection), SQLModel.metadata) == []
    migrate(fresh.app.state.settings.database_url)
    assert fresh.get("/health").status_code == 200


def test_force_password_rotation_and_cookie_session(fresh):
    assert fresh.get("/api/dashboard").status_code == 401
    response = fresh.post("/api/auth/login", json={"username": "admin", "password": TEST_PASSWORD})
    assert response.json()["must_change_password"]
    assert "HttpOnly" in response.headers.get_list("set-cookie")[0]
    assert "Max-Age=2592000" in response.headers.get_list("set-cookie")[0]
    assert fresh.get("/api/dashboard").status_code == 403
    fresh.headers["x-csrf-token"] = fresh.cookies["farm_csrf"]
    assert (
        fresh.post(
            "/api/auth/password",
            json={"current_password": TEST_PASSWORD, "new_password": NEW_PASSWORD},
        ).status_code
        == 200
    )
    assert fresh.get("/api/dashboard").status_code == 200
    assert fresh.post("/api/auth/logout").status_code == 200
    assert fresh.get("/api/auth/session").status_code == 401


def test_csrf_and_origin_are_required_for_mutations(client):
    assert client.post("/api/auth/logout", headers={"x-csrf-token": ""}).status_code == 403
    assert (
        client.post("/api/auth/logout", headers={"origin": "https://untrusted.example"}).status_code
        == 403
    )
    assert client.get("/api/auth/session").status_code == 200


def test_user_permissions_enforced_on_backend(client, opt):
    r = client.post("/api/users", json={"username": "worker", "password": "test-worker-password"})
    assert r.status_code == 201
    assert r.json()["role"] == "user"
    assert "password_hash" not in r.json()
    record = create(
        client, "production", {"date": str(now().date()), "group_id": opt(kind="group"), "kg": "1"}
    )
    assert client.post("/api/auth/logout").status_code == 200
    assert (
        client.post(
            "/api/auth/login", json={"username": "worker", "password": "test-worker-password"}
        ).status_code
        == 200
    )
    client.headers["x-csrf-token"] = client.cookies["farm_csrf"]
    assert (
        client.post(
            "/api/auth/password",
            json={
                "current_password": "test-worker-password",
                "new_password": "test-worker-password-new",
            },
        ).status_code
        == 200
    )
    assert client.get("/api/users").status_code == 403
    assert (
        client.post(
            "/api/records/options", json={"kind": "group", "name": "Unauthorized"}
        ).status_code
        == 403
    )
    assert client.delete(f"/api/records/production/{record['id']}").status_code == 403
    created = create(
        client, "production", {"date": str(now().date()), "group_id": opt(kind="group"), "kg": "2"}
    )
    assert created["created_by"] != record["created_by"]


def test_option_roles_and_references_survive_rename(client, opt, options):
    egg = next(o for o in options if o["role"] == "egg")
    data = {k: egg[k] for k in ["kind", "name", "unit", "tracks_quantity", "kg_per_sack"]}
    data["name"] = "Telur super"
    assert client.put(f"/api/records/options/{egg['id']}", json=data).status_code == 200
    created = create(client, "sales", sale(opt, quantity="2", unit_id=opt("box")))
    assert float(created["egg_kg"]) == 30
    assert created["product_name"] == "Telur super"
    assert client.delete(f"/api/records/options/{egg['id']}").status_code == 409
    group = create(client, "options", {"kind": "group", "name": "Kandang C"})
    create(client, "production", {"date": str(now().date()), "group_id": group["id"], "kg": "1"})
    assert client.delete(f"/api/records/options/{group['id']}").status_code == 409


def test_expense_preview_feed_factor_and_unit_snapshot(client, opt, options):
    data = {
        "date": str(now().date()),
        "category_id": opt("feed"),
        "feed_type_id": opt(kind="feed_type"),
        "quantity": "2.5",
        "unit_price": 300000,
    }
    response = client.post("/api/records/expenses/preview", json=data)
    assert response.status_code == 200
    assert response.json()["amount"] == 750000
    assert float(response.json()["kg"]) == 125
    expense = create(client, "expenses", data)
    category = next(o for o in options if o["role"] == "feed")
    update = {k: category[k] for k in ["kind", "name", "unit", "tracks_quantity", "kg_per_sack"]}
    update["unit"] = "Sak baru"
    assert client.put(f"/api/records/options/{category['id']}", json=update).status_code == 200
    changed = client.put(
        f"/api/records/expenses/{expense['id']}", json={**data, "notes": "Nota diperbarui"}
    )
    assert changed.status_code == 200
    assert changed.json()["unit"] == "Karung"
    assert float(changed.json()["kg"]) == 125


def test_sales_overpayment_edit_delete_and_cascade(client, opt):
    record = create(client, "sales", sale(opt, first_payment=260000))
    assert record["total"] == 250000
    assert record["overpaid"] == 10000
    assert record["status"] == "paid"
    detail = client.get(f"/api/records/sales/{record['id']}").json()
    payment_id = detail["history"][0]["id"]
    payment = {"date": record["date"], "amount": 100000}
    assert (
        client.put(
            f"/api/records/sales/{record['id']}/history/{payment_id}", json=payment
        ).status_code
        == 200
    )
    assert client.get(f"/api/records/sales/{record['id']}").json()["outstanding"] == 150000
    assert (
        client.delete(f"/api/records/sales/{record['id']}/history/{payment_id}").status_code == 200
    )
    assert client.get(f"/api/records/sales/{record['id']}").json()["status"] == "unpaid"
    assert (
        client.post(f"/api/records/sales/{record['id']}/history", json=payment).status_code == 201
    )
    assert client.delete(f"/api/records/sales/{record['id']}").status_code == 200
    assert client.get(f"/api/records/sales/{record['id']}").status_code == 404
    assert client.get("/api/dashboard").json()["receivables"] == 0


def test_debt_repayments_and_date_guards(client):
    day = now().date()
    record = create(
        client,
        "debts",
        {"date": str(day), "lender": "Koperasi", "amount": 1000000, "interest_rate": "5"},
    )
    path = f"/api/records/debts/{record['id']}/history"
    assert (
        client.post(path, json={"date": str(day - timedelta(days=1)), "amount": 1}).status_code
        == 422
    )
    assert client.post(path, json={"date": str(day), "amount": 1100000}).status_code == 201
    detail = client.get(f"/api/records/debts/{record['id']}").json()
    assert detail["outstanding"] == 0
    assert detail["overpaid"] == 100000
    assert detail["status"] == "paid"
    assert client.get("/api/records/debts?q=Koperasi&status=paid").json()["count"] == 1


def test_flock_limits_and_proportional_book_value(client, opt):
    data = {
        "date": str(now().date()),
        "name": "Layer A",
        "type_id": opt("chicken"),
        "group_id": opt(kind="group"),
        "quantity": "100",
        "unit_price": 50000,
        "depreciation_months": 24,
    }
    asset = create(client, "assets", data)
    assert asset["book_value"] == 5000000
    path = f"/api/records/assets/{asset['id']}/history"
    assert (
        client.post(
            path, json={"date": data["date"], "quantity": "10", "reason": "sold"}
        ).status_code
        == 201
    )
    changed = client.get(f"/api/records/assets/{asset['id']}").json()
    assert float(changed["remaining_quantity"]) == 90
    assert changed["book_value"] == 4500000
    assert (
        client.post(
            path, json={"date": data["date"], "quantity": "91", "reason": "dead"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            path, json={"date": data["date"], "quantity": "0.5", "reason": "missing"}
        ).status_code
        == 422
    )
    assert (
        client.put(
            f"/api/records/assets/{asset['id']}", json={**data, "quantity": "90"}
        ).status_code
        == 422
    )
    history_id = changed["history"][0]["id"]
    assert (
        client.put(f"{path}/{history_id}", json={"date": data["date"], "amount": 1}).status_code
        == 405
    )
    assert client.delete(f"{path}/{history_id}").status_code == 200
    assert client.get(f"/api/records/assets/{asset['id']}").json()["book_value"] == 5000000


def test_dashboard_stock_ratios_and_financial_export(client, opt):
    day = str(now().date())
    create(
        client,
        "assets",
        {
            "date": day,
            "name": "Flock",
            "type_id": opt("chicken"),
            "group_id": opt(kind="group"),
            "quantity": "100",
            "unit_price": 50000,
            "depreciation_months": 24,
        },
    )
    create(client, "production", {"date": day, "group_id": opt(kind="group"), "kg": "6"})
    create(
        client,
        "expenses",
        {"date": day, "category_id": opt("egg_purchase"), "quantity": "10", "unit_price": 20000},
    )
    create(
        client,
        "expenses",
        {
            "date": day,
            "category_id": opt("feed"),
            "feed_type_id": opt(kind="feed_type"),
            "quantity": "1",
            "unit_price": 300000,
        },
    )
    create(client, "sales", sale(opt, quantity="1", unit_id=opt("box"), paid_in_full=True))
    result = client.get("/api/dashboard?period=today").json()
    assert float(result["stock_kg"]) == 1
    assert float(result["production_kg"]) == 6
    assert result["estimated_eggs"] == 100
    assert float(result["hdp"]) == 100
    assert float(result["fcr"]) == pytest.approx(50 / 6)
    assert result["receivables"] == 0
    report = client.get("/api/reports/financial").json()
    assert report["month"]["revenue"] == 25000
    assert report["month"]["cost_of_goods"] == 500000
    assert report["month"]["net_profit"] == -475000
    response = client.get("/api/reports/export")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book["Laba Rugi"]["B3"].value == 25000
    assert book["Tren Bulanan"].max_row == 13


@pytest.mark.parametrize(
    "kind,data",
    [
        ("production", {"date": "bad", "group_id": 1, "kg": "-1"}),
        ("debts", {"date": "2026-01-01", "amount": 1.5, "lender": "X"}),
        ("options", {"kind": "group", "name": ""}),
    ],
)
def test_uniform_validation_errors(client, kind, data):
    response = client.post(f"/api/records/{kind}", json=data)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == 422
    assert response.json()["error"]["details"]


def test_invalid_foreign_option_and_market_failure(client):
    response = client.post(
        "/api/records/production", json={"date": str(now().date()), "group_id": 9999, "kg": "1"}
    )
    assert response.status_code == 422
    assert client.get("/api/market-prices").json()["status"] == "not_configured"
    assert client.post("/api/market-prices/refresh").status_code == 503
    assert (
        client.get("/api/dashboard?period=custom&start=2026-02-01&end=2026-01-01").status_code
        == 422
    )


def test_month_end_depreciation_matches_report(client, opt):
    from datetime import date

    from app.models import Asset
    from app.reports import pnl
    from app.rules import asset_view
    from sqlmodel import Session

    asset = create(
        client,
        "assets",
        {
            "date": "2026-01-31",
            "name": "Layer month end",
            "type_id": opt("chicken"),
            "group_id": opt(kind="group"),
            "quantity": "10",
            "unit_price": 120000,
            "depreciation_months": 12,
        },
    )
    with Session(client.app.state.engine) as db:
        record = db.get(Asset, asset["id"])
        february = asset_view(db, record, date(2026, 2, 28))
        assert february["elapsed_months"] == 1
        assert february["book_value"] == 1100000
        assert pnl(db, date(2026, 2, 1), date(2026, 2, 28))["depreciation"] == 100000
        assert asset_view(db, record, date(2027, 1, 31))["book_value"] == 0


def test_computed_money_limit_returns_validation_error(client, opt):
    response = client.post(
        "/api/records/sales", json=sale(opt, quantity="1000000000", unit_price=1000000000000)
    )
    assert response.status_code == 422
    assert response.json()["error"]["message"]


def test_concurrent_flock_reductions_are_serialized(client, opt):
    from concurrent.futures import ThreadPoolExecutor

    asset = create(
        client,
        "assets",
        {
            "date": str(now().date()),
            "name": "Concurrent flock",
            "type_id": opt("chicken"),
            "group_id": opt(kind="group"),
            "quantity": "100",
            "unit_price": 1000,
            "depreciation_months": 12,
        },
    )
    path = f"/api/records/assets/{asset['id']}/history"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda _: (
                    client.post(
                        path, json={"date": str(now().date()), "quantity": "60", "reason": "dead"}
                    ).status_code
                ),
                range(2),
            )
        )
    assert sorted(results) == [201, 422]
    detail = client.get(f"/api/records/assets/{asset['id']}").json()
    assert float(detail["remaining_quantity"]) == 40
